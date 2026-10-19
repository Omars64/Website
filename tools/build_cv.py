"""Build the shareable CV from the portfolio's current verified content."""

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "cv" / "Omar_Solanki_CV.docx"
OUT.parent.mkdir(exist_ok=True)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.48)
section.bottom_margin = Inches(0.45)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(9)
normal.font.color.rgb = RGBColor(35, 45, 58)
normal.paragraph_format.space_after = Pt(3)

for name, size, before, after in [
    ("Title", 20, 0, 3),
    ("Heading 1", 10.5, 7, 2),
    ("Heading 2", 9.5, 4, 1),
]:
    style = doc.styles[name]
    style.font.name = "Aptos"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
for border in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(border)

bullet_style = doc.styles.add_style("CV Bullet", WD_STYLE_TYPE.PARAGRAPH)
bullet_style.base_style = normal
bullet_style.paragraph_format.left_indent = Inches(0.18)
bullet_style.paragraph_format.first_line_indent = Inches(-0.1)
bullet_style.paragraph_format.space_after = Pt(2)

def line(text, style=None):
    return doc.add_paragraph(text, style)

def bullet(text):
    line("•  " + text, "CV Bullet")

title = doc.add_paragraph("Omar Solanki", "Title")
title.runs[0].font.name = "Aptos"
title.runs[0].font.size = Pt(20)
title.runs[0].font.bold = True
title.runs[0].font.color.rgb = RGBColor(0, 0, 0)
line("Junior Software Engineer | RPA and Automation Developer | Python and Process Automation")
line("Salmiya, Kuwait  |  omarsolanki35@gmail.com")

doc.add_heading("Professional Profile", 1)
line("Junior Software Engineer at Digital Processing Systems Kuwait with experience in Python development, enterprise automation, databases, and system integration. Works with Selenium, Blue Prism, Flowable, APIs, and web technologies, with growing hands-on experience in Gemini Enterprise and AI agents.")

doc.add_heading("Technical Skills", 1)
line("Python · Selenium · Blue Prism · RPA · Flowable · Liferay integration · webMethods API integration · SQL · Microsoft SQL Server · Oracle Database · JavaScript · Full-stack development · REST and Web Services APIs · Gemini Enterprise · Google ADK · AI agents")

doc.add_heading("Professional Experience", 1)
doc.add_heading("Junior Software Engineer | Digital Processing Systems Kuwait", 2)
line("August 2023 – Present  |  Salmiya, Kuwait")
doc.add_heading("Department of Legal Advice and Legislation (FATWA) | Dec 2023–Aug 2025", 2)
bullet("Developed business automation using Blue Prism and Python with Selenium.")
bullet("Managed a major part of the archiving operation, including its databases and systems.")
bullet("Monitored and supported the archiving workflow to keep daily operations running smoothly.")
doc.add_heading("DPS Internal Workspace | Aug–Nov 2025", 2)
bullet("Independently designed and built a full-stack internal workspace application for DPS.")
bullet("Built a Jira-like platform to support team work and project coordination.")
doc.add_heading("Flowable Workflow Integration | Dec 2025–Apr 2026", 2)
bullet("Served as the dedicated Flowable resource, responsible for the workflow engine, backend routing, and process flows.")
bullet("Integrated Flowable workflows with the Liferay user interface and coordinated with the Liferay frontend team.")
bullet("Worked with the webMethods team on API integration between the workflow engine and connected systems.")
doc.add_heading("Warba Insurance and Reinsurance Company | May 2026–Present", 2)
bullet("Develop and maintain Selenium automation for two motor lines of business, covering the business workflows and requirements of both lines.")
bullet("Coordinate with the team to keep automation behavior aligned with databases and related systems across those workflows.")

doc.add_heading("Earlier Experience", 2)
bullet("Sales Associate | SAI INFORMATICS | Jul–Aug 2023 | Doha, Qatar")
bullet("Higher Education Consultant | AlShamlan edu | Jun–Aug 2023 | Shuwaikh, Kuwait")
bullet("Data Entry Specialist and Sales Executive | Al braimi | Apr 2021–May 2023 | Amghara, Kuwait")

doc.add_heading("Education", 1)
bullet("Bachelor of Computer Applications, Computer Science | Jaipur National University | Jun 2020–Apr 2023")
bullet("High School Diploma, Computer Science | Board of School Education Hubli Karnataka | Jun 2017–Jun 2019")
bullet("Hifz e Quran | Masjid Ghanimah tul Muzaini | Apr 2016–May 2018")

doc.add_heading("Licenses and Certifications", 1)
certs = [
    "Deploy Gemini Enterprise with Workspace Data Sources and Model Armor | Google | Oct 2026",
    "Certified Partner Specialist Gemini Enterprise Deployment | Google | Oct 2026–Apr 2027",
    "Govern Agent Access with Gemini Enterprise Agent Platform | Google | Oct 2026",
    "Add Agents to Gemini Enterprise | Google | Sep 2026",
    "Deploy an Agent with Agent Development Kit (ADK) | Google | Sep 2026",
    "Accelerate Development with Antigravity | Google | Sep 2026",
    "The Complete Blue Prism Masterclass: Beginner to Expert | Udemy | Nov 2023",
]
for cert in certs:
    bullet(cert)

line("Languages: Arabic (Professional Working), English (Native or Bilingual)  |  Honor: Hafiz-e-Quran")

doc.save(OUT)
print(OUT)
